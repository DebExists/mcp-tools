import os
import tempfile
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from readability import Document
from ddgs import DDGS
from fastmcp import FastMCP

mcp = FastMCP("TeleClaw Tools")

MAX_DOWNLOAD_BYTES = 15 * 1024 * 1024
TIMEOUT = 20

def _check_http_url(url: str):
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise ValueError("Please provide a valid public http or https URL.")

def _download(url: str):
    _check_http_url(url)
    r = requests.get(
        url,
        timeout=TIMEOUT,
        stream=True,
        headers={"User-Agent": "TeleClawTools/1.0"},
    )
    r.raise_for_status()

    data = bytearray()
    for chunk in r.iter_content(chunk_size=65536):
        if chunk:
            data.extend(chunk)
        if len(data) > MAX_DOWNLOAD_BYTES:
            raise ValueError("File is larger than the 15 MB limit.")
    return bytes(data), r.headers.get("content-type", "")

@mcp.tool()
def web_search(query: str) -> str:
    """Search the public web. No API key is required."""
    results = DDGS().text(query, max_results=5)
    output = []
    for item in results:
        output.append(
            f"Title: {item.get('title', '')}\n"
            f"URL: {item.get('href', '')}\n"
            f"Content: {item.get('body', '')}"
        )
    return "\n\n---\n\n".join(output) or "No results found."

@mcp.tool()
def read_webpage(url: str) -> str:
    """Read the main text from a public webpage."""
    _check_http_url(url)
    r = requests.get(
        url,
        timeout=TIMEOUT,
        headers={"User-Agent": "TeleClawTools/1.0"},
    )
    r.raise_for_status()

    try:
        html = Document(r.text).summary()
    except Exception:
        html = r.text

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    return soup.get_text("\n", strip=True)[:100000]

@mcp.tool()
def read_file(file_url: str) -> str:
    """Read a public PDF, DOCX, TXT, MD, CSV, or XLSX file from a URL."""
    data, content_type = _download(file_url)
    suffix = Path(urlparse(file_url).path).suffix.lower()

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
        f.write(data)
        path = f.name

    try:
        if suffix == ".pdf" or "pdf" in content_type:
            import fitz
            doc = fitz.open(path)
            return "\n".join(page.get_text() for page in doc)[:100000]

        if suffix == ".docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(path)
            return "\n".join(p.text for p in doc.paragraphs)[:100000]

        if suffix in (".txt", ".md", ".json"):
            return data.decode("utf-8", errors="replace")[:100000]

        if suffix == ".csv":
            import pandas as pd
            return pd.read_csv(path).to_csv(index=False)[:100000]

        if suffix == ".xlsx":
            import pandas as pd
            sheets = pd.read_excel(path, sheet_name=None)
            return "\n\n".join(
                f"Sheet: {name}\n{df.to_csv(index=False)}"
                for name, df in sheets.items()
            )[:100000]

        return "Unsupported file type. Supported: PDF, DOCX, TXT, MD, CSV, XLSX."
    finally:
        try:
            os.remove(path)
        except OSError:
            pass

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8000"))
    mcp.run(transport="http", host="0.0.0.0", port=port)
